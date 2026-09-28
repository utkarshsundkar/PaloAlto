"""
Security Policies List Definition
Contains structured security policy definitions and example vHEADER policy definitions.
Supports parsing and converting Cisco Firepower / vHEADER schema rules to Palo Alto SecurityPolicy format.
"""

from dataclasses import dataclass
from typing import Any
from excel_reader import SecurityPolicy


# ─────────────────────────────────────────────────────────────
# vHEADER Definition & Example Policies (Pipe-Delimited)
# ─────────────────────────────────────────────────────────────
vHEADER = (
    "device.name | rule.number | rule.name | rule.disabled | rule.sourceZone | "
    "rule.source | rule.destinationZone | rule.destination | rule.service | "
    "rule.action | rule.actionbehavior | rule.securityProfile | rule.hitcount | "
    "rule.lastused | rule.cumulativeseverity | rule.changeddate | rule.changeUser | "
    "rule.identifier | rule.createDate | rule.direction | rule.comment | rule.fmName"
)

# Canonical alias
VHEADER = vHEADER

EXAMPLE_POLICIES_TEXT = """
FPR2140-MUM-1 | 1 | Allow Internal Web Traffic | false | Inside_zone | 10.10.10.0/24 | Inside_zone,DMZ_zone | 10.20.20.10 | HTTP,HTTPS | ACCEPT | DECISIVELY | | 25 | | 0 | 2026-09-28T10:00:00.000Z | admin | 11111111-1111-1111-1111-111111111111 | 2026-09-28T10:00:00.000Z | NONE | Allow internal users to web server | FPR2140-MUM-1

FPR2140-MUM-1 | 2 | Allow DNS Queries | false | Inside_zone | 10.10.10.0/24 | Outside_zone | Any | DNS | ACCEPT | DECISIVELY | | 120 | | 0 | 2026-09-28T10:05:00.000Z | admin | 22222222-2222-2222-2222-222222222222 | 2026-09-28T10:05:00.000Z | NONE | Allow DNS resolution | FPR2140-MUM-1

FPR2140-MUM-1 | 3 | Allow HTTPS Internet Access | false | Inside_zone | 10.10.20.0/24 | Outside_zone | Any | HTTPS | ACCEPT | DECISIVELY | | 350 | | 0 | 2026-09-28T10:10:00.000Z | admin | 33333333-3333-3333-3333-333333333333 | 2026-09-28T10:10:00.000Z | NONE | Allow secure internet access | FPR2140-MUM-1

FPR2140-MUM-1 | 4 | Allow SSH To Management Server | false | Inside_zone | 10.10.30.0/24 | Inside_zone | 10.10.30.10 | SSH | ACCEPT | DECISIVELY | | 15 | | 0 | 2026-09-28T10:15:00.000Z | admin | 44444444-4444-4444-4444-444444444444 | 2026-09-28T10:15:00.000Z | NONE | Allow administrators to management server | FPR2140-MUM-1

FPR2140-MUM-1 | 5 | Allow Application Server To Database | false | DMZ_zone | 10.20.20.0/24 | Inside_zone | 10.30.30.20 | TCP-3306 | ACCEPT | DECISIVELY | | 42 | | 0 | 2026-09-28T10:20:00.000Z | admin | 55555555-5555-5555-5555-555555555555 | 2026-09-28T10:20:00.000Z | NONE | Application database connectivity | FPR2140-MUM-1

FPR2140-MUM-1 | 6 | Allow VPN Users | false | VPN_zone | Any | Inside_zone | 10.10.40.0/24 | Any | ACCEPT | DECISIVELY | | 87 | | 0 | 2026-09-28T10:25:00.000Z | admin | 66666666-6666-6666-6666-666666666666 | 2026-09-28T10:25:00.000Z | NONE | Allow VPN users to internal network | FPR2140-MUM-1

FPR2140-MUM-1 | 7 | Allow NTP Traffic | false | Inside_zone | 10.10.10.0/24 | Outside_zone | Any | NTP | ACCEPT | DECISIVELY | | 31 | | 0 | 2026-09-28T10:30:00.000Z | admin | 77777777-7777-7777-7777-777777777777 | 2026-09-28T10:30:00.000Z | NONE | Allow time synchronization | FPR2140-MUM-1

FPR2140-MUM-1 | 8 | Allow Mail Server Traffic | false | Inside_zone | 10.10.50.0/24 | DMZ_zone | 10.20.20.30 | SMTP | ACCEPT | DECISIVELY | | 18 | | 0 | 2026-09-28T10:35:00.000Z | admin | 88888888-8888-8888-8888-888888888888 | 2026-09-28T10:35:00.000Z | NONE | Allow mail server communication | FPR2140-MUM-1

FPR2140-MUM-1 | 9 | Block Guest Network To Internal | false | Guest_zone | 10.99.0.0/24 | Inside_zone | Any | Any | DROP | DECISIVELY | | 0 | | 0 | 2026-09-28T10:40:00.000Z | admin | 99999999-9999-9999-9999-999999999999 | 2026-09-28T10:40:00.000Z | NONE | Prevent guest network access to internal resources | FPR2140-MUM-1

FPR2140-MUM-1 | 10 | Allow Monitoring Server | false | Inside_zone | 10.10.60.10 | DMZ_zone | 10.20.20.0/24 | ICMP | ACCEPT | DECISIVELY | | 9 | | 0 | 2026-09-28T10:45:00.000Z | admin | aaaaaaaa-aaaa-aaaa-aaaa-aaaaaaaaaaaa | 2026-09-28T10:45:00.000Z | NONE | Allow monitoring server connectivity checks | FPR2140-MUM-1
"""


def parse_vheader_policies(raw_text: str = EXAMPLE_POLICIES_TEXT, header_line: str = vHEADER) -> list[dict[str, str]]:
    """
    Parses pipe-delimited text containing policies into a list of dictionaries,
    keyed by the column names from vHEADER.
    """
    headers = [h.strip() for h in header_line.split("|")]
    records: list[dict[str, str]] = []

    for line in raw_text.strip().splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        cols = [c.strip() for c in line.split("|")]
        # Skip line if it's the header itself
        if cols[0] == headers[0] or "rule.name" in line:
            continue
        if len(cols) == len(headers):
            records.append(dict(zip(headers, cols)))
        elif len(cols) > 0:
            # Pad with empty strings if column count mismatch
            padded = cols + [""] * (len(headers) - len(cols))
            records.append(dict(zip(headers, padded[:len(headers)])))

    return records


# Parsed list of raw vHEADER policy dictionaries
VHEADER_EXAMPLE_POLICIES: list[dict[str, str]] = parse_vheader_policies(EXAMPLE_POLICIES_TEXT, vHEADER)


def map_service_to_panos(service: str) -> tuple[str, str]:
    """
    Map Cisco/vHEADER service names to PAN-OS application + service.
    Returns (application, service).
    """
    raw = (service or "any").strip()
    if not raw or raw.lower() == "any":
        return "any", "any"

    apps: list[str] = []
    svcs: list[str] = []
    for part in [p.strip() for p in raw.split(",") if p.strip()]:
        key = part.lower()
        if key in ("http",):
            apps.append("web-browsing")
            svcs.append("application-default")
        elif key in ("https", "ssl"):
            apps.append("ssl")
            svcs.append("application-default")
        elif key in ("dns",):
            apps.append("dns")
            svcs.append("application-default")
        elif key in ("ssh",):
            apps.append("ssh")
            svcs.append("application-default")
        elif key in ("ntp",):
            apps.append("ntp")
            svcs.append("application-default")
        elif key in ("smtp",):
            apps.append("smtp")
            svcs.append("application-default")
        elif key in ("icmp", "ping"):
            apps.append("ping")
            svcs.append("application-default")
        elif key.startswith("tcp-") and key[4:].isdigit():
            svcs.append(part.upper() if part.upper().startswith("TCP-") else f"TCP-{key[4:]}")
            apps.append("any")
        else:
            svcs.append(part)
            apps.append("any")

    unique_apps = []
    for a in apps:
        if a not in unique_apps:
            unique_apps.append(a)
    unique_svcs = []
    for s in svcs:
        if s not in unique_svcs:
            unique_svcs.append(s)
    if set(unique_svcs) == {"application-default"}:
        return ", ".join(unique_apps), "application-default"
    if "application-default" in unique_svcs and len(unique_svcs) > 1:
        unique_svcs = [s for s in unique_svcs if s != "application-default"]
    return ", ".join(unique_apps) if unique_apps else "any", ", ".join(unique_svcs) if unique_svcs else "any"


def vheader_to_security_policy_dict(item: dict[str, str]) -> dict[str, Any]:
    """
    Translates a vHEADER rule dictionary to Palo Alto SecurityPolicy definition dict.
    Maps:
      - ACCEPT -> allow, DROP -> deny
      - Any -> any
      - Source / Destination Zones, Addresses, and Services
    """
    action_raw = item.get("rule.action", "ACCEPT").strip().upper()
    action = "allow" if action_raw in ("ACCEPT", "ALLOW", "PERMIT") else "deny"

    # Normalize 'Any' to 'any' for consistency
    src_addr = item.get("rule.source", "any").strip()
    if src_addr.lower() == "any":
        src_addr = "any"

    dst_addr = item.get("rule.destination", "any").strip()
    if dst_addr.lower() == "any":
        dst_addr = "any"

    service = item.get("rule.service", "any").strip()
    if service.lower() == "any":
        service = "any"

    is_disabled = item.get("rule.disabled", "false").strip().lower() == "true"

    application, pan_service = map_service_to_panos(service)

    return {
        "name": item.get("rule.name", "").strip(),
        "srczone": item.get("rule.sourceZone", "any").strip() or "any",
        "dstzone": item.get("rule.destinationZone", "any").strip() or "any",
        "srcaddr": src_addr or "any",
        "dstaddr": dst_addr or "any",
        "application": application,
        "service": pan_service,
        "action": action,
        "profile_group": item.get("rule.securityProfile", "").strip(),
        "log_start": "no",
        "log_end": "yes",
        "description": item.get("rule.comment", "").strip(),
        "tag": item.get("rule.fmName", "").strip() or item.get("device.name", "").strip(),
        "disabled": "yes" if is_disabled else "no",
        "hitcount": item.get("rule.hitcount", "").strip(),
        "rule_number": item.get("rule.number", "").strip(),
        "identifier": item.get("rule.identifier", "").strip(),
    }


# ─────────────────────────────────────────────────────────────
# Converted vHEADER Security Policies Data (Primary list)
# ─────────────────────────────────────────────────────────────
SECURITY_POLICIES_DATA = [
    vheader_to_security_policy_dict(p) for p in VHEADER_EXAMPLE_POLICIES
]


# ─────────────────────────────────────────────────────────────
# Original Baseline Enterprise Security Policies (Preserved)
# ─────────────────────────────────────────────────────────────
BASELINE_SECURITY_POLICIES_DATA = [
    {
        "name": "MGMT-ADMIN-ACCESS",
        "srczone": "TRUST",
        "dstzone": "TRUST",
        "srcaddr": "CORP-LAN",
        "dstaddr": "any",
        "application": "ssh, ssl, ping",
        "service": "application-default",
        "action": "allow",
        "profile_group": "strict",
        "log_start": "no",
        "log_end": "yes",
        "description": "Administrative and management access from secure LAN",
        "tag": "management",
    },
    {
        "name": "INFRA-DNS-NTP",
        "srczone": "TRUST",
        "dstzone": "UNTRUST",
        "srcaddr": "CORP-LAN",
        "dstaddr": "any",
        "application": "dns, ntp",
        "service": "application-default",
        "action": "allow",
        "profile_group": "strict",
        "log_start": "no",
        "log_end": "yes",
        "description": "Outbound DNS and NTP time sync infrastructure services",
        "tag": "infrastructure",
    },
    {
        "name": "CORP-INTERNET-OUT",
        "srczone": "TRUST",
        "dstzone": "UNTRUST",
        "srcaddr": "CORP-LAN",
        "dstaddr": "any",
        "application": "web-browsing, ssl",
        "service": "application-default",
        "action": "allow",
        "profile_group": "strict",
        "log_start": "no",
        "log_end": "yes",
        "description": "Standard outbound corporate web access with inspection",
        "tag": "corp",
    },
    {
        "name": "DMZ-INBOUND-HTTPS",
        "srczone": "UNTRUST",
        "dstzone": "DMZ",
        "srcaddr": "any",
        "dstaddr": "DMZ-HOST",
        "application": "web-browsing, ssl",
        "service": "SVC-HTTPS-8443",
        "action": "allow",
        "profile_group": "strict",
        "log_start": "no",
        "log_end": "yes",
        "description": "Inbound public traffic to DMZ HTTPS service",
        "tag": "dmz",
    },
    {
        "name": "DMZ-TO-INTERNAL-DB",
        "srczone": "DMZ",
        "dstzone": "TRUST",
        "srcaddr": "DMZ-HOST",
        "dstaddr": "CORP-LAN",
        "application": "ms-sql-db, mysql, postgresql",
        "service": "application-default",
        "action": "allow",
        "profile_group": "strict",
        "log_start": "no",
        "log_end": "yes",
        "description": "Database tier connectivity from DMZ application hosts",
        "tag": "dmz-db",
    },
    {
        "name": "VPN-TO-INTERNAL-TRUST",
        "srczone": "VPN-ZONE",
        "dstzone": "TRUST",
        "srcaddr": "any",
        "dstaddr": "CORP-LAN",
        "application": "web-browsing, ssl, ssh, ms-rdp",
        "service": "application-default",
        "action": "allow",
        "profile_group": "strict",
        "log_start": "no",
        "log_end": "yes",
        "description": "Remote VPN client access to internal network resources",
        "tag": "vpn",
    },
    {
        "name": "BLOCK-GUEST-TO-TRUST",
        "srczone": "guest",
        "dstzone": "TRUST",
        "srcaddr": "GUEST-RANGE",
        "dstaddr": "CORP-LAN",
        "application": "any",
        "service": "any",
        "action": "deny",
        "profile_group": "",
        "log_start": "no",
        "log_end": "yes",
        "description": "Isolate guest network from corporate internal trust zone",
        "tag": "security-isolation",
    },
    {
        "name": "BLOCK-INSECURE-PROTOCOLS",
        "srczone": "any",
        "dstzone": "any",
        "srcaddr": "any",
        "dstaddr": "any",
        "application": "ms-ds-smb, netbios-ssn, telnet",
        "service": "application-default",
        "action": "deny",
        "profile_group": "",
        "log_start": "no",
        "log_end": "yes",
        "description": "Block vulnerable cleartext and SMB protocols across all zones",
        "tag": "hardening",
    },
    {
        "name": "DEFAULT-DENY-ALL",
        "srczone": "any",
        "dstzone": "any",
        "srcaddr": "any",
        "dstaddr": "any",
        "application": "any",
        "service": "any",
        "action": "deny",
        "profile_group": "",
        "log_start": "no",
        "log_end": "yes",
        "description": "Zero-trust explicit default deny baseline with logging",
        "tag": "baseline",
    },
]


def get_security_policies(use_baseline: bool = False) -> list[SecurityPolicy]:
    """
    Return the list of SecurityPolicy dataclass instances.
    By default returns the 10 vHEADER example policies.
    Pass use_baseline=True to return the baseline enterprise policies.
    """
    dataset = BASELINE_SECURITY_POLICIES_DATA if use_baseline else SECURITY_POLICIES_DATA
    policies = []
    for idx, item in enumerate(dataset, start=2):
        policies.append(SecurityPolicy(
            row_index=idx,
            name=item["name"],
            srczone=item.get("srczone", "any"),
            dstzone=item.get("dstzone", "any"),
            srcaddr=item.get("srcaddr", "any"),
            dstaddr=item.get("dstaddr", "any"),
            application=item.get("application", "any"),
            service=item.get("service", "any"),
            action=item.get("action", "allow"),
            profile_group=item.get("profile_group", ""),
            log_start=item.get("log_start", "no"),
            log_end=item.get("log_end", "yes"),
            description=item.get("description", ""),
            tag=item.get("tag", ""),
            disabled=item.get("disabled", "no"),
        ))
    return policies


def export_policies_to_excel(excel_path: str = "policies.xlsx", use_baseline: bool = False) -> int:
    """
    Exports/syncs the security policy definitions directly into the
    Security_Policies sheet of the target Excel workbook.
    """
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    import config

    EX_FILL = PatternFill("solid", fgColor="EBF3FB")
    EX_FONT = Font(italic=False, color="000000", size=10)
    THIN = Side(border_style="thin", color="BFBFBF")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    MID = Alignment(vertical="center")

    wb = load_workbook(excel_path)
    sheet_name = getattr(config, "SHEET_SECURITY_POLICIES", "Security_Policies")
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(sheet_name)
    else:
        ws = wb[sheet_name]

    # Clear old rows below header
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    policies = get_security_policies(use_baseline=use_baseline)
    for row_idx, pol in enumerate(policies, start=2):
        row_vals = [
            pol.name,
            pol.srczone,
            pol.dstzone,
            pol.srcaddr,
            pol.dstaddr,
            pol.application,
            pol.service,
            pol.action,
            pol.profile_group,
            pol.log_start,
            pol.log_end,
            pol.description,
            pol.tag,
            "",  # Status
            "",  # Error Detail
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = EX_FILL
            cell.font = EX_FONT
            cell.border = BORDER
            cell.alignment = MID

    wb.save(excel_path)
    return len(policies)
