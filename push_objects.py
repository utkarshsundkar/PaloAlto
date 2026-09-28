"""
Palo Alto Networks Object Configuration Pusher
Configures enterprise Address Objects and Service Objects tailored to the network
topology, pushes them to the firewall, commits, and synchronizes policies.xlsx.
"""

import sys
import logging
from rich.console import Console
from rich.table import Table
from rich.panel import Panel
from rich import box
import openpyxl

import config
from ssh_client import PaloAltoSSHClient
from cli_builder import CLIBuilder
from excel_reader import AddressObject, ServiceObject
from policy_pusher import _has_error

if sys.platform == "win32":
    try:
        if hasattr(sys.stdout, "reconfigure"):
            sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        if hasattr(sys.stderr, "reconfigure"):
            sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

console = Console()
log = logging.getLogger("paloalto.objects")

# ─────────────────────────────────────────────────────────────
#  Enterprise Address Objects
# ─────────────────────────────────────────────────────────────
ADDRESS_OBJECTS = [
    # Enterprise Subnets
    AddressObject(
        row_index=2,
        name="NET-LAN-Users",
        type="ip-netmask",
        value="10.10.10.0/24",
        description="Corporate Users & Workstations Subnet",
        tag="FPR2140-MUM-1",
    ),
    AddressObject(
        row_index=3,
        name="NET-LAN-Servers",
        type="ip-netmask",
        value="10.10.20.0/24",
        description="Internal Corporate Servers Subnet",
        tag="FPR2140-MUM-1",
    ),
    AddressObject(
        row_index=4,
        name="NET-LAN-Mgmt",
        type="ip-netmask",
        value="10.10.30.0/24",
        description="Network and IT Infrastructure Management Subnet",
        tag="FPR2140-MUM-1",
    ),
    AddressObject(
        row_index=5,
        name="NET-VPN-Pool",
        type="ip-netmask",
        value="10.10.40.0/24",
        description="Remote Access VPN Client IP Pool",
        tag="FPR2140-MUM-1",
    ),
    AddressObject(
        row_index=6,
        name="NET-Mail-Clients",
        type="ip-netmask",
        value="10.10.50.0/24",
        description="Internal Mail Clients and SMTP Relays",
        tag="FPR2140-MUM-1",
    ),
    AddressObject(
        row_index=7,
        name="NET-DMZ-Servers",
        type="ip-netmask",
        value="10.20.20.0/24",
        description="Public-Facing DMZ Server Subnet",
        tag="FPR2140-MUM-1",
    ),
    AddressObject(
        row_index=8,
        name="NET-Guest-WiFi",
        type="ip-netmask",
        value="10.99.0.0/24",
        description="Isolated Guest Wi-Fi Network",
        tag="FPR2140-MUM-1",
    ),
    AddressObject(
        row_index=9,
        name="NET-Database-Farm",
        type="ip-netmask",
        value="10.30.30.0/24",
        description="Backend Database Farm Subnet",
        tag="FPR2140-MUM-1",
    ),

    # Specific Server Hosts
    AddressObject(
        row_index=10,
        name="HOST-DMZ-Web",
        type="ip-netmask",
        value="10.20.20.10/32",
        description="Production DMZ Web Application Server",
        tag="FPR2140-MUM-1",
    ),
    AddressObject(
        row_index=11,
        name="HOST-DMZ-Mail",
        type="ip-netmask",
        value="10.20.20.30/32",
        description="Production DMZ Mail and SMTP Server",
        tag="FPR2140-MUM-1",
    ),
    AddressObject(
        row_index=12,
        name="HOST-DB-Server",
        type="ip-netmask",
        value="10.30.30.20/32",
        description="Production Database Server (MySQL/Oracle)",
        tag="FPR2140-MUM-1",
    ),
    AddressObject(
        row_index=13,
        name="HOST-Mgmt-Server",
        type="ip-netmask",
        value="10.10.30.10/32",
        description="Central IT Bastion / Jump Server",
        tag="FPR2140-MUM-1",
    ),
    AddressObject(
        row_index=14,
        name="HOST-NMS-Monitoring",
        type="ip-netmask",
        value="10.10.60.10/32",
        description="Enterprise NMS / Monitoring Host",
        tag="FPR2140-MUM-1",
    ),

    # Public External VIPs & Pool IPs
    AddressObject(
        row_index=15,
        name="PUB-Web-Server-VIP",
        type="ip-netmask",
        value="203.0.113.20/32",
        description="External Public VIP for Inbound Web Traffic",
        tag="FPR2140-MUM-1",
    ),
    AddressObject(
        row_index=16,
        name="PUB-Mail-Server-VIP",
        type="ip-netmask",
        value="203.0.113.25/32",
        description="External Public VIP for Inbound Mail Traffic",
        tag="FPR2140-MUM-1",
    ),
    AddressObject(
        row_index=17,
        name="PUB-PAT-Internet-IP",
        type="ip-netmask",
        value="203.0.113.10/32",
        description="Public Egress NAT Pool for Corporate Users",
        tag="FPR2140-MUM-1",
    ),

    # FQDN and IP Range
    AddressObject(
        row_index=18,
        name="FQDN-PaloAlto-Updates",
        type="fqdn",
        value="updates.paloaltonetworks.com",
        description="Palo Alto Dynamic Content Updates FQDN",
        tag="FPR2140-MUM-1",
    ),
    AddressObject(
        row_index=19,
        name="RNG-DHCP-Pool",
        type="ip-range",
        value="10.10.10.100-10.10.10.200",
        description="Corporate DHCP Client Allocation Pool",
        tag="FPR2140-MUM-1",
    ),
]

# ─────────────────────────────────────────────────────────────
#  Enterprise Service Objects
# ─────────────────────────────────────────────────────────────
SERVICE_OBJECTS = [
    ServiceObject(
        row_index=2,
        name="TCP-3306",
        protocol="tcp",
        dst_port="3306",
        description="MySQL / MariaDB Database Service",
        tag="FPR2140-MUM-1",
    ),
    ServiceObject(
        row_index=3,
        name="SVC-HTTPS-8443",
        protocol="tcp",
        dst_port="8443",
        description="Custom Web Administration HTTPS",
        tag="FPR2140-MUM-1",
    ),
    ServiceObject(
        row_index=4,
        name="SVC-HTTP-8080",
        protocol="tcp",
        dst_port="8080",
        description="Alternative HTTP Web Proxy / App Server",
        tag="FPR2140-MUM-1",
    ),
    ServiceObject(
        row_index=5,
        name="SVC-RDP-3389",
        protocol="tcp",
        dst_port="3389",
        description="Microsoft Remote Desktop Protocol",
        tag="FPR2140-MUM-1",
    ),
    ServiceObject(
        row_index=6,
        name="SVC-SSH-2222",
        protocol="tcp",
        dst_port="2222",
        description="Alternate SSH Administration Port",
        tag="FPR2140-MUM-1",
    ),
    ServiceObject(
        row_index=7,
        name="SVC-MSSQL-1433",
        protocol="tcp",
        dst_port="1433",
        description="Microsoft SQL Server Database Port",
        tag="FPR2140-MUM-1",
    ),
    ServiceObject(
        row_index=8,
        name="SVC-SYSLOG-UDP-514",
        protocol="udp",
        dst_port="514",
        description="Syslog / SIEM UDP Log Receiver",
        tag="FPR2140-MUM-1",
    ),
    ServiceObject(
        row_index=9,
        name="SVC-SNMP-UDP-161",
        protocol="udp",
        dst_port="161",
        description="SNMP Polling and Monitoring Service",
        tag="FPR2140-MUM-1",
    ),
    ServiceObject(
        row_index=10,
        name="SVC-NTP-UDP-123",
        protocol="udp",
        dst_port="123",
        description="Network Time Protocol Synchronization",
        tag="FPR2140-MUM-1",
    ),
    ServiceObject(
        row_index=11,
        name="SVC-DNS-UDP-53",
        protocol="udp",
        dst_port="53",
        description="Domain Name System Standard UDP Queries",
        tag="FPR2140-MUM-1",
    ),
]


def sync_objects_to_excel(excel_path: str = config.EXCEL_FILE) -> None:
    """Updates policies.xlsx Address_Objects and Service_Objects sheets."""
    try:
        wb = openpyxl.load_workbook(excel_path)

        # 1. Sync Address Objects
        sheet_addr = getattr(config, "SHEET_ADDRESS_OBJECTS", "Address_Objects")
        if sheet_addr in wb.sheetnames:
            ws_addr = wb[sheet_addr]
            if ws_addr.max_row > 1:
                ws_addr.delete_rows(2, ws_addr.max_row - 1)
        else:
            ws_addr = wb.create_sheet(sheet_addr)
            ws_addr.append(["Name", "Type", "Value", "Description", "Tag(s)", "Status", "Error Detail"])

        for row_idx, addr in enumerate(ADDRESS_OBJECTS, start=2):
            row_data = [
                addr.name,
                addr.type,
                addr.value,
                addr.description,
                addr.tag,
                addr.status,
                addr.error_msg,
            ]
            for col_idx, val in enumerate(row_data, start=1):
                ws_addr.cell(row=row_idx, column=col_idx, value=val)

        # 2. Sync Service Objects
        sheet_svc = getattr(config, "SHEET_SERVICE_OBJECTS", "Service_Objects")
        if sheet_svc in wb.sheetnames:
            ws_svc = wb[sheet_svc]
            if ws_svc.max_row > 1:
                ws_svc.delete_rows(2, ws_svc.max_row - 1)
        else:
            ws_svc = wb.create_sheet(sheet_svc)
            ws_svc.append(["Name", "Protocol", "Dst Port", "Src Port", "Description", "Tag(s)", "Status", "Error Detail"])

        for row_idx, svc in enumerate(SERVICE_OBJECTS, start=2):
            row_data = [
                svc.name,
                svc.protocol,
                svc.dst_port,
                svc.src_port or "",
                svc.description,
                svc.tag,
                svc.status,
                svc.error_msg,
            ]
            for col_idx, val in enumerate(row_data, start=1):
                ws_svc.cell(row=row_idx, column=col_idx, value=val)

        wb.save(excel_path)
        console.print(f"[dim]Synced {len(ADDRESS_OBJECTS)} address objects and {len(SERVICE_OBJECTS)} service objects into {excel_path}[/]")
    except Exception as e:
        console.print(f"[yellow]Warning: Could not update Excel sheets: {e}[/]")


def push_objects() -> int:
    console.print(Panel.fit(
        "[bold cyan]Palo Alto Object Configuration (Address & Service)[/]\n"
        f"[dim]Target Firewall   : [white]{config.PALOALTO_HOST}:{config.PALOALTO_PORT}[/white]\n"
        f"Address Objects   : [white]{len(ADDRESS_OBJECTS)}[/white]\n"
        f"Service Objects   : [white]{len(SERVICE_OBJECTS)}[/white][/]",
        border_style="cyan",
        title="[bold white]Objects Deployment[/]",
    ))

    # 1. Address Objects Preview Table
    addr_tbl = Table(title="Target Address Objects", box=box.ROUNDED, header_style="bold cyan", show_lines=True)
    addr_tbl.add_column("#", style="dim", width=4)
    addr_tbl.add_column("Object Name", style="bold white")
    addr_tbl.add_column("Type", style="yellow")
    addr_tbl.add_column("Value / Subnet", style="green")
    addr_tbl.add_column("Description", style="dim")
    addr_tbl.add_column("Status", no_wrap=True)

    for idx, addr in enumerate(ADDRESS_OBJECTS, start=1):
        addr_tbl.add_row(str(idx), addr.name, addr.type, addr.value, addr.description, "[yellow]PENDING[/]")

    console.print(addr_tbl)

    # 2. Service Objects Preview Table
    svc_tbl = Table(title="Target Service Objects", box=box.ROUNDED, header_style="bold cyan", show_lines=True)
    svc_tbl.add_column("#", style="dim", width=4)
    svc_tbl.add_column("Service Name", style="bold white")
    svc_tbl.add_column("Protocol", style="yellow")
    svc_tbl.add_column("Dst Port", style="green")
    svc_tbl.add_column("Description", style="dim")
    svc_tbl.add_column("Status", no_wrap=True)

    for idx, svc in enumerate(SERVICE_OBJECTS, start=1):
        svc_tbl.add_row(str(idx), svc.name, svc.protocol.upper(), svc.dst_port, svc.description, "[yellow]PENDING[/]")

    console.print(svc_tbl)

    builder = CLIBuilder()

    console.print(f"\nConnecting to firewall [bold white]{config.PALOALTO_HOST}[/] ...")
    ssh = PaloAltoSSHClient(
        host=config.PALOALTO_HOST,
        port=config.PALOALTO_PORT,
        username=config.PALOALTO_USERNAME,
        password=config.PALOALTO_PASSWORD,
        timeout=config.SSH_TIMEOUT,
        cmd_timeout=config.CMD_TIMEOUT,
        banner_timeout=config.BANNER_TIMEOUT,
    )
    ssh.connect()
    console.print("   [green]Connected successfully.[/]")

    ssh.configure()
    del_base = f"delete vsys {config.VSYS}" if getattr(config, "USE_VSYS_PREFIX", False) else "delete"

    # Push Address Objects
    console.print(f"\n[bold]Pushing {len(ADDRESS_OBJECTS)} Address Objects ...[/]\n")
    for idx, addr in enumerate(ADDRESS_OBJECTS, start=1):
        name = addr.name
        console.print(f" [{idx}/{len(ADDRESS_OBJECTS)}] Processing Address: [bold cyan]{name}[/] ({addr.type}: {addr.value})")

        # Delete existing object to avoid conflict
        ssh.send_cfg_cmd(f'{del_base} address "{name}"')

        cmds = builder.address_object(addr)
        failed = False
        for cmd in cmds:
            out = ssh.send_cfg_cmd(cmd)
            err, msg = _has_error(out)
            if err:
                console.print(f"       [bold red][FAIL][/] {msg}")
                addr.status = "FAILED"
                addr.error_msg = msg
                failed = True
                break

        if not failed:
            addr.status = "PUSHED"
            console.print(f"       [green][OK] Success[/] -> {name}")

    # Push Service Objects
    console.print(f"\n[bold]Pushing {len(SERVICE_OBJECTS)} Service Objects ...[/]\n")
    for idx, svc in enumerate(SERVICE_OBJECTS, start=1):
        name = svc.name
        console.print(f" [{idx}/{len(SERVICE_OBJECTS)}] Processing Service: [bold cyan]{name}[/] ({svc.protocol.upper()}/{svc.dst_port})")

        # Delete existing object to avoid conflict
        ssh.send_cfg_cmd(f'{del_base} service "{name}"')

        cmds = builder.service_object(svc)
        failed = False
        for cmd in cmds:
            out = ssh.send_cfg_cmd(cmd)
            err, msg = _has_error(out)
            if err:
                console.print(f"       [bold red][FAIL][/] {msg}")
                svc.status = "FAILED"
                svc.error_msg = msg
                failed = True
                break

        if not failed:
            svc.status = "PUSHED"
            console.print(f"       [green][OK] Success[/] -> {name}")

    # Commit
    console.print("\n[bold]Committing Object configuration on firewall ...[/]")
    commit_out = ssh.commit(description="Configured enterprise Address and Service Objects")
    console.print("   [green]Commit completed successfully![/]")

    ssh.disconnect()

    # Sync to Excel
    sync_objects_to_excel()

    # Final summary tables
    res_addr_tbl = Table(title="Address Objects Push Results", box=box.ROUNDED, header_style="bold cyan", show_lines=True)
    res_addr_tbl.add_column("#", style="dim", width=4)
    res_addr_tbl.add_column("Object Name", style="bold white")
    res_addr_tbl.add_column("Type", style="yellow")
    res_addr_tbl.add_column("Value / Subnet", style="green")
    res_addr_tbl.add_column("Status", no_wrap=True)

    for idx, addr in enumerate(ADDRESS_OBJECTS, start=1):
        st = "[bold green]OK PUSHED[/]" if addr.status == "PUSHED" else "[bold red]FAIL[/]"
        res_addr_tbl.add_row(str(idx), addr.name, addr.type, addr.value, st)

    console.print("\n")
    console.print(res_addr_tbl)

    res_svc_tbl = Table(title="Service Objects Push Results", box=box.ROUNDED, header_style="bold cyan", show_lines=True)
    res_svc_tbl.add_column("#", style="dim", width=4)
    res_svc_tbl.add_column("Service Name", style="bold white")
    res_svc_tbl.add_column("Protocol / Port", style="cyan")
    res_svc_tbl.add_column("Status", no_wrap=True)

    for idx, svc in enumerate(SERVICE_OBJECTS, start=1):
        st = "[bold green]OK PUSHED[/]" if svc.status == "PUSHED" else "[bold red]FAIL[/]"
        res_svc_tbl.add_row(str(idx), svc.name, f"{svc.protocol.upper()}/{svc.dst_port}", st)

    console.print("\n")
    console.print(res_svc_tbl)

    addr_pushed = sum(1 for a in ADDRESS_OBJECTS if a.status == "PUSHED")
    addr_failed = sum(1 for a in ADDRESS_OBJECTS if a.status == "FAILED")
    svc_pushed = sum(1 for s in SERVICE_OBJECTS if s.status == "PUSHED")
    svc_failed = sum(1 for s in SERVICE_OBJECTS if s.status == "FAILED")

    console.print(
        f"\n[bold]Objects Deployment Complete:[/]\n"
        f"  Address Objects : [green]{addr_pushed} Pushed[/] | [red]{addr_failed} Failed[/] | Total: {len(ADDRESS_OBJECTS)}\n"
        f"  Service Objects : [green]{svc_pushed} Pushed[/] | [red]{svc_failed} Failed[/] | Total: {len(SERVICE_OBJECTS)}\n"
    )

    return 1 if (addr_failed + svc_failed) > 0 else 0


if __name__ == "__main__":
    sys.exit(push_objects())
