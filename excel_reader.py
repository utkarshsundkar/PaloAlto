"""
Excel Reader
Reads all policy sheets from the Excel workbook and returns typed dataclasses.
"""

import logging
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

import config

log = logging.getLogger("paloalto.excel")


# ──────────────────────────────────────────────────────
#  Data classes — one per PAN-OS object type
# ──────────────────────────────────────────────────────

@dataclass
class AddressObject:
    row_index: int
    name: str
    type: str           # ip-netmask | ip-range | fqdn | ip-wildcard
    value: str          # CIDR / range "start-end" / FQDN string / wildcard
    description: str = ""
    tag: str = ""
    status: str = "PENDING"
    error_msg: str = ""


@dataclass
class ServiceObject:
    row_index: int
    name: str
    protocol: str       # tcp | udp | tcp-udp
    dst_port: str       # e.g. 8443, 8080-8090
    src_port: str = ""
    description: str = ""
    tag: str = ""
    status: str = "PENDING"
    error_msg: str = ""


@dataclass
class SecurityZone:
    row_index: int
    name: str
    mode: str           # layer3 | layer2 | virtual-wire | tap | tunnel
    interfaces: str     # comma-separated: ethernet1/1, ethernet1/2
    enable_userid: str = "no"
    description: str = ""
    status: str = "PENDING"
    error_msg: str = ""


@dataclass
class NATPolicy:
    row_index: int
    name: str
    srczone: str        # source zone(s) comma-sep
    dstzone: str        # destination zone
    srcaddr: str        # source address object(s) comma-sep
    dstaddr: str        # destination address object(s)
    service: str        # "any" or service object name
    nat_type: str       # ipv4 | nat64
    snat_type: str      # dynamic-ip-and-port | dynamic-ip | static-ip | none
    snat_interface: str = ""    # egress interface for DIPP
    snat_ip: str = ""           # static/dynamic translated IP
    dnat_address: str = ""      # destination NAT translated address
    dnat_port: str = ""         # destination NAT translated port
    description: str = ""
    tag: str = ""
    status: str = "PENDING"
    error_msg: str = ""


@dataclass
class SecurityPolicy:
    row_index: int
    name: str
    srczone: str        # comma-sep zone names
    dstzone: str
    srcaddr: str        # comma-sep address objects
    dstaddr: str
    application: str    # comma-sep: any, web-browsing, ssl …
    service: str        # application-default | any | service-object
    action: str         # allow | deny | drop | reset-client | reset-server | reset-both
    profile_group: str = ""     # security profile group
    log_start: str = "no"
    log_end: str = "yes"
    description: str = ""
    tag: str = ""
    status: str = "PENDING"
    error_msg: str = ""
    disabled: str = "no"


@dataclass
class PolicyWorkbook:
    address_objects:   list[AddressObject]  = field(default_factory=list)
    service_objects:   list[ServiceObject]  = field(default_factory=list)
    security_zones:    list[SecurityZone]   = field(default_factory=list)
    nat_policies:      list[NATPolicy]      = field(default_factory=list)
    security_policies: list[SecurityPolicy] = field(default_factory=list)


# ──────────────────────────────────────────────────────
#  Reader
# ──────────────────────────────────────────────────────

class ExcelReader:
    def __init__(self, filepath: str = config.EXCEL_FILE):
        self.filepath = Path(filepath)
        if not self.filepath.exists():
            raise FileNotFoundError(f"Excel file not found: {self.filepath.resolve()}")
        self.wb: Workbook = openpyxl.load_workbook(self.filepath, data_only=True)

    def read_all(self) -> PolicyWorkbook:
        pwb = PolicyWorkbook()
        pwb.address_objects   = self._read_address_objects()
        pwb.service_objects   = self._read_service_objects()
        pwb.security_zones    = self._read_security_zones()
        pwb.nat_policies      = self._read_nat_policies()
        pwb.security_policies = self._read_security_policies()
        log.info(
            f"Loaded: {len(pwb.address_objects)} addr | "
            f"{len(pwb.service_objects)} svc | "
            f"{len(pwb.security_zones)} zones | "
            f"{len(pwb.nat_policies)} NAT | "
            f"{len(pwb.security_policies)} security policies"
        )
        return pwb

    # ──────────────────────────────────────
    #  Helpers
    # ──────────────────────────────────────
    def _sheet(self, name: str) -> Worksheet | None:
        if name not in self.wb.sheetnames:
            log.warning(f"Sheet '{name}' not found — skipping.")
            return None
        return self.wb[name]

    @staticmethod
    def _v(row, col: int) -> str:
        val = row[col - 1].value
        return str(val).strip() if val is not None else ""

    # ──────────────────────────────────────
    #  Per-sheet readers
    # ──────────────────────────────────────
    def _read_address_objects(self) -> list[AddressObject]:
        ws = self._sheet(config.SHEET_ADDRESS_OBJECTS)
        if not ws:
            return []
        result = []
        for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
            name = self._v(row, 1)
            if not name or name.startswith("#"):
                continue
            result.append(AddressObject(
                row_index=i,
                name=name,
                type=self._v(row, 2).lower() or "ip-netmask",
                value=self._v(row, 3),
                description=self._v(row, 4),
                tag=self._v(row, 5),
            ))
        return result

    def _read_service_objects(self) -> list[ServiceObject]:
        ws = self._sheet(config.SHEET_SERVICE_OBJECTS)
        if not ws:
            return []
        result = []
        for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
            name = self._v(row, 1)
            if not name or name.startswith("#"):
                continue
            result.append(ServiceObject(
                row_index=i,
                name=name,
                protocol=self._v(row, 2).lower() or "tcp",
                dst_port=self._v(row, 3),
                src_port=self._v(row, 4),
                description=self._v(row, 5),
                tag=self._v(row, 6),
            ))
        return result

    def _read_security_zones(self) -> list[SecurityZone]:
        ws = self._sheet(config.SHEET_SECURITY_ZONES)
        if not ws:
            return []
        result = []
        for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
            name = self._v(row, 1)
            if not name or name.startswith("#"):
                continue
            result.append(SecurityZone(
                row_index=i,
                name=name,
                mode=self._v(row, 2).lower() or "layer3",
                interfaces=self._v(row, 3),
                enable_userid=self._v(row, 4).lower() or "no",
                description=self._v(row, 5),
            ))
        return result

    def _read_nat_policies(self) -> list[NATPolicy]:
        ws = self._sheet(config.SHEET_NAT_POLICIES)
        if not ws:
            return []
        result = []
        for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
            name = self._v(row, 1)
            if not name or name.startswith("#"):
                continue
            result.append(NATPolicy(
                row_index=i,
                name=name,
                srczone=self._v(row, 2),
                dstzone=self._v(row, 3),
                srcaddr=self._v(row, 4),
                dstaddr=self._v(row, 5),
                service=self._v(row, 6) or "any",
                nat_type=self._v(row, 7).lower() or "ipv4",
                snat_type=self._v(row, 8).lower() or "dynamic-ip-and-port",
                snat_interface=self._v(row, 9),
                snat_ip=self._v(row, 10),
                dnat_address=self._v(row, 11),
                dnat_port=self._v(row, 12),
                description=self._v(row, 13),
                tag=self._v(row, 14),
            ))
        return result

    def _read_security_policies(self) -> list[SecurityPolicy]:
        ws = self._sheet(config.SHEET_SECURITY_POLICIES)
        if not ws:
            return []
        result = []
        for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
            name = self._v(row, 1)
            if not name or name.startswith("#"):
                continue
            result.append(SecurityPolicy(
                row_index=i,
                name=name,
                srczone=self._v(row, 2),
                dstzone=self._v(row, 3),
                srcaddr=self._v(row, 4),
                dstaddr=self._v(row, 5),
                application=self._v(row, 6) or "any",
                service=self._v(row, 7) or "application-default",
                action=self._v(row, 8).lower() or "allow",
                profile_group=self._v(row, 9),
                log_start=self._v(row, 10).lower() or "no",
                log_end=self._v(row, 11).lower() or "yes",
                description=self._v(row, 12),
                tag=self._v(row, 13),
            ))
        return result
