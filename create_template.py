"""
Template Generator
Run this ONCE to create a ready-to-fill policies.xlsx with correct
headers, example rows, and column widths for all 5 sheets.
"""

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import config

# ── Styles ────────────────────────────────────────────
HDR_FILL = PatternFill("solid", fgColor="1F3864")    # dark navy
EX_FILL  = PatternFill("solid", fgColor="EBF3FB")    # light blue
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
EX_FONT  = Font(italic=True, color="444444", size=10)
THIN     = Side(border_style="thin", color="BFBFBF")
BORDER   = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER   = Alignment(horizontal="center", vertical="center", wrap_text=True)
MID      = Alignment(vertical="center")


def _header(ws, headers: list[str], widths: list[int]) -> None:
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = HDR_FILL; c.font = HDR_FONT
        c.alignment = CENTER; c.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def _row(ws, r: int, vals: list) -> None:
    for col, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=col, value=v)
        c.fill = EX_FILL; c.font = EX_FONT
        c.border = BORDER; c.alignment = MID


def create_template(out: str = config.EXCEL_FILE) -> None:
    wb = Workbook()

    # ── 1. Address Objects ────────────────────────────
    ws = wb.active
    ws.title = config.SHEET_ADDRESS_OBJECTS
    _header(ws,
        ["Name", "Type", "Value", "Description", "Tag(s)", "Status", "Error Detail"],
        [28, 14, 28, 35, 20, 12, 40])
    for i, ex in enumerate([
        ["CORP-LAN",    "ip-netmask", "192.168.10.0/24",  "Corporate LAN subnet",   "corp"],
        ["DMZ-HOST",    "ip-netmask", "10.0.0.50/32",     "DMZ web server",         "dmz"],
        ["GUEST-RANGE", "ip-range",   "172.16.0.1-172.16.0.254", "Guest pool",      "guest"],
        ["EXT-API",     "fqdn",       "api.example.com",  "External API endpoint",  "external"],
    ], start=2):
        _row(ws, i, ex)

    # ── 2. Service Objects ────────────────────────────
    ws2 = wb.create_sheet(config.SHEET_SERVICE_OBJECTS)
    _header(ws2,
        ["Name", "Protocol", "Dst Port", "Src Port", "Description", "Tag(s)", "Status", "Error Detail"],
        [28, 12, 16, 16, 35, 20, 12, 40])
    for i, ex in enumerate([
        ["SVC-HTTPS-8443", "tcp",     "8443",  "",           "Custom HTTPS",  "web"],
        ["SVC-DNS-UDP",    "udp",     "53",    "",           "DNS",           "infra"],
        ["SVC-RDP",        "tcp",     "3389",  "",           "Remote Desktop","mgmt"],
        ["SVC-HTTP-HTTPS", "tcp-udp", "80,443","1024-65535", "Web combo",     "web"],
    ], start=2):
        _row(ws2, i, ex)

    # ── 3. Security Zones ─────────────────────────────
    ws3 = wb.create_sheet(config.SHEET_SECURITY_ZONES)
    _header(ws3,
        ["Name", "Mode", "Interfaces (comma-sep)", "Enable User-ID", "Description", "Status", "Error Detail"],
        [22, 14, 35, 18, 35, 12, 40])
    for i, ex in enumerate([
        ["TRUST",     "layer3",       "ethernet1/1",        "yes", "Internal trust zone"],
        ["UNTRUST",   "layer3",       "ethernet1/2",        "no",  "Internet-facing zone"],
        ["DMZ",       "layer3",       "ethernet1/3",        "no",  "DMZ zone"],
        ["VPN-ZONE",  "tunnel",       "",                   "no",  "VPN tunnel zone"],
    ], start=2):
        _row(ws3, i, ex)

    # ── 4. NAT Policies ──────────────────────────────
    ws4 = wb.create_sheet(config.SHEET_NAT_POLICIES)
    _header(ws4,
        ["Name", "Src Zone", "Dst Zone", "Src Address", "Dst Address",
         "Service", "NAT Type", "SNAT Type", "SNAT Interface",
         "SNAT IP", "DNAT Address", "DNAT Port",
         "Description", "Tag(s)", "Status", "Error Detail"],
        [28,14,14,22,22,14,10,22,18,18,18,12,35,18,12,40])
    for i, ex in enumerate([
        ["TRUST-TO-INTERNET-NAT", "TRUST", "UNTRUST", "CORP-LAN", "any",
         "any", "ipv4", "dynamic-ip-and-port", "ethernet1/2", "", "", "",
         "Outbound NAT for corp", "nat"],
        ["INBOUND-DMZ-DNAT", "UNTRUST", "DMZ", "any", "any",
         "SVC-HTTPS-8443", "ipv4", "none", "", "", "10.0.0.50", "8443",
         "Port-forward to DMZ server", "dnat"],
    ], start=2):
        _row(ws4, i, ex)

    # ── 5. Security Policies ──────────────────────────
    ws5 = wb.create_sheet(config.SHEET_SECURITY_POLICIES)
    _header(ws5,
        ["Name", "Src Zone", "Dst Zone", "Src Address", "Dst Address",
         "Application", "Service", "Action",
         "Profile Group", "Log Start", "Log End",
         "Description", "Tag(s)", "Status", "Error Detail"],
        [30,14,14,22,22,25,22,10,20,10,10,35,18,12,40])
    for i, ex in enumerate([
        ["CORP-INTERNET-OUT", "TRUST", "UNTRUST", "CORP-LAN", "any",
         "web-browsing, ssl", "application-default", "allow",
         "strict", "no", "yes", "Corp internet access", "corp"],
        ["DMZ-INBOUND", "UNTRUST", "DMZ", "any", "DMZ-HOST",
         "any", "SVC-HTTPS-8443", "allow",
         "strict", "no", "yes", "Public to DMZ HTTPS", "dmz"],
        ["BLOCK-GUEST-TO-TRUST", "guest", "TRUST", "GUEST-RANGE", "CORP-LAN",
         "any", "any", "deny",
         "", "no", "yes", "Block guest from corp", "security"],
        ["DEFAULT-DENY", "any", "any", "any", "any",
         "any", "any", "deny",
         "", "no", "yes", "Explicit default deny", "baseline"],
    ], start=2):
        _row(ws5, i, ex)

    wb.save(out)
    print(f"\n[OK] Template created: {out}")
    print("  Fill in your data, then run:  python main.py")
    print("  Rows whose Name starts with '#' are skipped.\n")


if __name__ == "__main__":
    create_template()
