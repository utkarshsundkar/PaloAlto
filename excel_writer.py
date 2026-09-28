"""
Excel Status Writer
Writes PUSHED / FAILED / DRY-RUN status back into the Excel workbook
with color-coded cells.
"""

import logging
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment

import config
from excel_reader import PolicyWorkbook

log = logging.getLogger("paloalto.excel_writer")

# ── Colour fills ──────────────────────────────────────
GREEN  = PatternFill("solid", fgColor="C6EFCE")   # pushed OK
RED    = PatternFill("solid", fgColor="FFC7CE")   # failed
YELLOW = PatternFill("solid", fgColor="FFEB9C")   # dry-run / pending
BLUE   = PatternFill("solid", fgColor="BDD7EE")   # header

# Column index of the Status and Error columns per sheet
# (1-based; must match the column layout in create_template.py)
_STATUS_COL: dict[str, int] = {
    config.SHEET_ADDRESS_OBJECTS:   6,
    config.SHEET_SERVICE_OBJECTS:   7,
    config.SHEET_SECURITY_ZONES:    6,
    config.SHEET_NAT_POLICIES:      15,
    config.SHEET_SECURITY_POLICIES: 14,
}


class ExcelWriter:
    def __init__(self, filepath: str = config.EXCEL_FILE):
        self.filepath = Path(filepath)

    def write_results(self, pwb: PolicyWorkbook) -> None:
        """Open the workbook in place and stamp every row with its result."""
        log.info(f"Writing results back to {self.filepath} …")
        wb = load_workbook(self.filepath)

        self._stamp(wb, config.SHEET_ADDRESS_OBJECTS,   pwb.address_objects)
        self._stamp(wb, config.SHEET_SERVICE_OBJECTS,   pwb.service_objects)
        self._stamp(wb, config.SHEET_SECURITY_ZONES,    pwb.security_zones)
        self._stamp(wb, config.SHEET_NAT_POLICIES,      pwb.nat_policies)
        self._stamp(wb, config.SHEET_SECURITY_POLICIES, pwb.security_policies)

        wb.save(self.filepath)
        log.info("Excel file updated ✔")

    # ──────────────────────────────────────────
    def _stamp(self, wb, sheet_name: str, items: list) -> None:
        if sheet_name not in wb.sheetnames or not items:
            return

        ws   = wb[sheet_name]
        scol = _STATUS_COL[sheet_name]
        ecol = scol + 1

        # Ensure header labels exist
        ws.cell(row=1, column=scol).value = "Status"
        ws.cell(row=1, column=ecol).value = "Error / Detail"
        ws.cell(row=1, column=scol).fill  = BLUE
        ws.cell(row=1, column=ecol).fill  = BLUE

        for item in items:
            r = item.row_index

            sc = ws.cell(row=r, column=scol)
            ec = ws.cell(row=r, column=ecol)

            sc.value = item.status
            ec.value = item.error_msg or ""

            fill = (
                GREEN  if item.status == "PUSHED"  else
                RED    if item.status == "FAILED"  else
                YELLOW
            )
            sc.fill  = fill
            sc.font  = Font(bold=True)
            sc.alignment = Alignment(horizontal="center")
