"""
Populates policies.xlsx from the parsed Cisco FTD Running Config.
Includes Address Objects, Service Objects, Security Zones, NAT Policies, and Security Policies.
"""

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import config
from parse_ftd_config import FTDConfigParser, sanitize_name

# Styles
HDR_FILL = PatternFill("solid", fgColor="1F3864")    # dark navy
ROW_FILL = PatternFill("solid", fgColor="F2F5F9")    # subtle light blue
ALT_FILL = PatternFill("solid", fgColor="FFFFFF")    # white
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(color="222222", size=10)
THIN = Side(border_style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
MID = Alignment(vertical="center")

def apply_header(ws, headers: list[str], widths: list[int]):
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = CENTER
        c.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

def populate_workbook(parser: FTDConfigParser, output_file: str = config.EXCEL_FILE):
    wb = openpyxl.Workbook()

    # ─────────────────────────────────────────────────────────
    # 1. Address Objects
    # ─────────────────────────────────────────────────────────
    ws_addr = wb.active
    ws_addr.title = config.SHEET_ADDRESS_OBJECTS
    addr_headers = ["Name", "Type", "Value", "Description", "Tag(s)", "Status", "Error Detail"]
    addr_widths  = [32, 16, 32, 40, 20, 12, 35]
    apply_header(ws_addr, addr_headers, addr_widths)

    r = 2
    for obj in parser.address_objects.values():
        fill = ROW_FILL if r % 2 == 0 else ALT_FILL
        vals = [obj.name, obj.obj_type, obj.value, obj.description, "FTD-MIGRATED", "PENDING", ""]
        for col, v in enumerate(vals, 1):
            c = ws_addr.cell(row=r, column=col, value=v)
            c.fill = fill
            c.font = BODY_FONT
            c.border = BORDER
            c.alignment = MID
        r += 1

    # ─────────────────────────────────────────────────────────
    # 2. Service Objects
    # ─────────────────────────────────────────────────────────
    ws_svc = wb.create_sheet(config.SHEET_SERVICE_OBJECTS)
    svc_headers = ["Name", "Protocol", "Dst Port", "Src Port", "Description", "Tag(s)", "Status", "Error Detail"]
    svc_widths  = [28, 12, 20, 14, 35, 20, 12, 35]
    apply_header(ws_svc, svc_headers, svc_widths)

    r = 2
    for svc in parser.service_objects.values():
        fill = ROW_FILL if r % 2 == 0 else ALT_FILL
        vals = [svc.name, svc.protocol, svc.port, "", svc.description, "FTD-MIGRATED", "PENDING", ""]
        for col, v in enumerate(vals, 1):
            c = ws_svc.cell(row=r, column=col, value=v)
            c.fill = fill
            c.font = BODY_FONT
            c.border = BORDER
            c.alignment = MID
        r += 1

    # ─────────────────────────────────────────────────────────
    # 3. Security Zones
    # ─────────────────────────────────────────────────────────
    ws_zone = wb.create_sheet(config.SHEET_SECURITY_ZONES)
    zone_headers = ["Name", "Mode", "Interfaces (comma-sep)", "Enable User-ID", "Description", "Status", "Error Detail"]
    zone_widths  = [26, 14, 30, 16, 35, 12, 35]
    apply_header(ws_zone, zone_headers, zone_widths)

    # Map zones to interfaces
    zone_ifaces = {}
    for iface in parser.interfaces.values():
        zone_ifaces.setdefault(iface.nameif, []).append(iface.interface.lower())

    r = 2
    for z in sorted(parser.zones):
        fill = ROW_FILL if r % 2 == 0 else ALT_FILL
        ifaces_str = ", ".join(zone_ifaces.get(z, []))
        vals = [z, "layer3", ifaces_str, "no", f"Migrated from Cisco FTD {z}", "PENDING", ""]
        for col, v in enumerate(vals, 1):
            c = ws_zone.cell(row=r, column=col, value=v)
            c.fill = fill
            c.font = BODY_FONT
            c.border = BORDER
            c.alignment = MID
        r += 1

    # ─────────────────────────────────────────────────────────
    # 4. NAT Policies
    # ─────────────────────────────────────────────────────────
    ws_nat = wb.create_sheet(config.SHEET_NAT_POLICIES)
    nat_headers = [
        "Name", "Src Zone", "Dst Zone", "Src Address", "Dst Address", "Service",
        "NAT Type", "SNAT Type", "SNAT Interface", "SNAT IP", "DNAT Address", "DNAT Port",
        "Description", "Tag(s)", "Status", "Error Detail"
    ]
    nat_widths = [32, 20, 20, 28, 28, 12, 10, 20, 16, 25, 25, 12, 35, 16, 12, 35]
    apply_header(ws_nat, nat_headers, nat_widths)

    r = 2
    for nr in parser.nat_rules:
        fill = ROW_FILL if r % 2 == 0 else ALT_FILL
        vals = [
            nr.name,
            nr.src_zone or "any",
            nr.dst_zone or "any",
            nr.src_addr or "any",
            nr.dst_addr or "any",
            nr.service or "any",
            "ipv4",
            nr.snat_type,
            nr.snat_interface,
            nr.trans_src if nr.snat_type != "none" else "",
            nr.trans_dst,
            "",
            nr.description,
            "FTD-MIGRATED",
            "PENDING",
            ""
        ]
        for col, v in enumerate(vals, 1):
            c = ws_nat.cell(row=r, column=col, value=v)
            c.fill = fill
            c.font = BODY_FONT
            c.border = BORDER
            c.alignment = MID
        r += 1

    # ─────────────────────────────────────────────────────────
    # 5. Security Policies
    # ─────────────────────────────────────────────────────────
    ws_sec = wb.create_sheet(config.SHEET_SECURITY_POLICIES)
    sec_headers = [
        "Name", "Src Zone", "Dst Zone", "Src Address", "Dst Address",
        "Application", "Service", "Action", "Profile Group", "Log Start", "Log End",
        "Description", "Tag(s)", "Status", "Error / Detail"
    ]
    sec_widths = [35, 20, 20, 30, 30, 15, 25, 12, 16, 12, 12, 40, 16, 12, 35]
    apply_header(ws_sec, sec_headers, sec_widths)

    r = 2
    for sr in parser.security_rules:
        fill = ROW_FILL if r % 2 == 0 else ALT_FILL
        vals = [
            sr.name,
            ", ".join(sr.src_zones),
            ", ".join(sr.dst_zones),
            ", ".join(sr.src_addrs),
            ", ".join(sr.dst_addrs),
            "any",
            ", ".join(sr.services) if sr.services else "any",
            sr.action,
            "",
            "no",
            sr.log_end,
            sr.description,
            "FTD-MIGRATED",
            "PENDING",
            ""
        ]
        for col, v in enumerate(vals, 1):
            c = ws_sec.cell(row=r, column=col, value=v)
            c.fill = fill
            c.font = BODY_FONT
            c.border = BORDER
            c.alignment = MID
        r += 1

    wb.save(output_file)
    print(f"Successfully generated {output_file} with all FTD migrated policies.")

if __name__ == "__main__":
    ftd_file = "Cisco FTD running config"
    parser = FTDConfigParser(ftd_file).parse()
    populate_workbook(parser, config.EXCEL_FILE)
